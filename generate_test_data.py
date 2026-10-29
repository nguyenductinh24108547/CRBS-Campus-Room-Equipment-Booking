"""Generate runtime test input JSON by joining Excel Test Cases with structured data definitions.

Excel is the source for case metadata and Test Data IDs. The structured setup/input/expected
parameters live in test_data/test_case_catalog.json because the workbook's Test Cases sheet
currently describes many of them in prose rather than executable fields.
Run from package root: python scripts/generate_test_data.py
"""
from __future__ import annotations

import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "CRBS_Requirements_Master.xlsx"
CATALOG = ROOT / "test_data" / "test_case_catalog.json"
TEMPLATE = ROOT / "test_data" / "acceptance_test_data_template.json"
OUTPUT = ROOT / "test_data" / "acceptance_test_data.json"


def records_by_id(ws, id_column: str) -> dict[str, dict]:
    headers = [cell.value for cell in ws[1]]
    if id_column not in headers:
        raise ValueError(f"Column {id_column!r} not found in sheet {ws.title!r}")
    id_idx = headers.index(id_column)
    result = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        key = row[id_idx]
        if key is None or str(key).strip() == "":
            continue
        key = str(key).strip()
        if key in result:
            raise ValueError(f"Duplicate {id_column} {key} in sheet {ws.title}")
        result[key] = {headers[i]: row[i] for i in range(len(headers))}
    return result


def clean(value):
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def main() -> None:
    if not WORKBOOK.exists():
        raise FileNotFoundError(f"Workbook not found: {WORKBOOK}")
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    if "test_cases" not in workbook.sheetnames or "test_data" not in workbook.sheetnames:
        raise ValueError("Workbook must contain 'test_cases' and 'test_data' sheets")
    excel_cases = records_by_id(workbook["test_cases"], "Test Case ID")
    excel_data = records_by_id(workbook["test_data"], "Test Data ID")
    workbook.close()

    catalog = json.loads(CATALOG.read_text(encoding="utf-8"))
    template = json.loads(TEMPLATE.read_text(encoding="utf-8"))
    generated_cases = {}
    for entry in catalog["test_cases"]:
        tc = entry["test_case"]
        data = entry["test_data"]["input_parameters"]
        if not data.get("execution_ready"):
            continue
        tc_id = tc["test_case_id"]
        data_id = data["test_data_id"]
        if tc_id not in excel_cases:
            raise ValueError(f"Executable case {tc_id} is missing from Excel sheet test_cases")
        if data_id not in excel_data:
            raise ValueError(f"Test Data ID {data_id} is missing from Excel sheet test_data")
        excel_case = excel_cases[tc_id]
        excel_data_row = excel_data[data_id]
        if str(excel_data_row.get("Test Case ID", "")).strip() != tc_id:
            raise ValueError(f"Excel Test Data ID {data_id} does not map to {tc_id}")

        # The Excel workbook supplies the traceable case metadata; the catalogue supplies
        # executable setup/input/expected parameters that are not yet separate Excel columns.
        generated_cases[tc_id] = {
            "data_id": data_id,
            "title": clean(excel_case.get("Title")) or tc.get("title"),
            "purpose": data["purpose"],
            "excel_metadata": {
                "type": clean(excel_case.get("Type")),
                "requirement_ids": clean(excel_case.get("Requirement IDs")),
                "use_case_ids": clean(excel_case.get("Use Case IDs")),
                "priority": clean(excel_case.get("Priority")),
                "preconditions": clean(excel_case.get("Preconditions")),
                "steps": clean(excel_case.get("Steps")),
                "expected_result": clean(excel_case.get("Expected Result")),
                "test_data_id": data_id,
                "test_data_source": clean(excel_data_row.get("Input Data File / JSON Path")),
            },
            "setup": data["setup"],
            "input": data["input"],
            "expected": data["expected"],
            "assertion": data["assertion"],
            "definition_status": data["definition_status"],
            "execution_ready": True,
            "scope": data.get("scope", "Educational prototype only"),
        }
        if data.get("review_note"):
            generated_cases[tc_id]["review_note"] = data["review_note"]

    if not generated_cases:
        raise ValueError("No execution-ready test cases were found")
    output = {
        "schema_version": "1.1",
        "project": template.get("project", "CRBS"),
        "scope": template.get("scope", "Educational prototype only"),
        "source_workbook": WORKBOOK.name,
        "source_sheets": ["test_cases", "test_data"],
        "generator": "scripts/generate_test_data.py",
        "base_fixture": template["base_fixture"],
        "cases": generated_cases,
        "execution_note": (
            "Generated by joining Excel case metadata/Test Data IDs with structured executable "
            "parameters in test_case_catalog.json. The workbook's prose-only fields are not by "
            "themselves executable input. Only execution_ready cases are emitted. Prototype only."
        ),
    }
    OUTPUT.write_text(json.dumps(output, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Generated {len(generated_cases)} executable cases -> {OUTPUT.relative_to(ROOT)}")
    print("Source: Excel test_cases/test_data metadata + structured parameter definitions in test_case_catalog.json")
    print("Scope: educational Python/SQLite prototype; not the separate target application")


if __name__ == "__main__":
    main()
